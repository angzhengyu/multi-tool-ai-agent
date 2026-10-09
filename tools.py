# -*- coding: utf-8 -*-
"""
Agent 的工具集
==============
工具清单（与 main.py 的 TOOLS_SCHEMA 保持一致）：
  get_current_time   取当前时间
  calculate          算数表达式（有字符白名单）
  list_files         列出项目里的文件
  describe_csv       在本地探查 CSV 结构（列名/类型/缺失/数值列统计，**不返回明细行**）
  read_excel         读 Excel 结构 + 双通道内容（公式文本 / 缓存值 / 缓存缺失统计），**只读**
  aggregate_csv      在本地做分组聚合（分组/求和/平均/计数/排序取前 N，**只返回聚合结果**）
  run_python         执行 Python 代码（可用 pandas / numpy）做数据分析
                     ↑ 能力最强也最危险，**默认关闭**，需 ENABLE_RUN_PYTHON=1 开启
"""
import os
import re
import subprocess
import sys
from datetime import datetime
from pathlib import Path

BASE_DIR = Path(__file__).parent
DATA_DIR = BASE_DIR / "data"


# ---------------- 简单工具 ----------------
def get_current_time() -> str:
    """工具：返回当前的日期和时间。"""
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def calculate(expression: str) -> str:
    """工具：计算一个数学表达式，例如 "12*35+7"。

    安全说明：只允许数字和 + - * / ( ) . % 这些字符，并用空的 __builtins__ 执行，
    避免"注入"（用户输入被当成代码执行）。
    """
    allowed = set("0123456789+-*/(). %")
    if not set(expression) <= allowed:
        return "表达式里含有不允许的字符（只支持数字和 + - * / ( ) . % ）"
    try:
        value = eval(expression, {"__builtins__": {}}, {})   # noqa: S307（已用字符白名单限制）
        return f"{expression} = {value}"
    except Exception as e:
        return f"计算出错：{type(e).__name__}: {e}"


# ---------------- 文件工具（限制在项目目录内，防止读到系统文件）----------------
def _safe_path(path_str: str) -> Path:
    p = (BASE_DIR / path_str).resolve()
    if p != BASE_DIR and BASE_DIR not in p.parents:
        raise ValueError("拒绝访问：只能操作项目目录内的文件")
    return p


def list_files(dir_path: str = ".") -> str:
    """工具：列出项目目录下的文件与文件夹。"""
    try:
        base = _safe_path(dir_path)
    except ValueError as e:
        return str(e)
    if not base.exists():
        return f"路径不存在：{dir_path}"
    if base.is_file():
        # 模型有时会把"文件"当"目录"传进来，这里给个明确提示，别让它崩
        return (f"「{dir_path}」是一个文件，不是目录。"
                f"请传入目录路径（例如 '.' 表示项目根目录，或 'data'）。")
    lines = []
    for p in sorted(base.iterdir()):
        if p.name in {"__pycache__", ".git", "chroma_db"}:
            continue
        lines.append(p.name + ("/" if p.is_dir() else ""))
    return "\n".join(lines) if lines else "(空目录)"


# ---------------- Excel 工作簿工具（财务场景）----------------
# 为什么要有这个专工具，而不是让模型用 run_python 裸写 openpyxl：
#   ① 读一张表要写二十几行代码、还要写两遍（公式 / 缓存值），实测吃掉 2~3 轮，
#      而主循环只有有限轮次；
#   ② 让模型自己回忆 openpyxl 用法，它会写出平台错误的路径（实测出现过 /home/user）；
#   ③ 最关键的一条：**「缓存值是否存在」必须由工具明确告知**，而不是让模型自己发现 ——
#      程序生成的 xlsx 缓存值全为 None，模型若把 None 当 0 继续算，整份报告就废了。
#   它仍然在沙箱内、只读项目目录，**没有放宽任何边界**。
_EXCEL_MAX_LIST = 40          # 每张表最多列多少个单元格（防止大表把输出灌爆）


def read_excel(path: str) -> str:
    """工具：读取一个 Excel 工作簿的**结构 + 双通道内容**，不返回批量明细。

    双通道（财务勾稽的必要条件）：
      · 通道 A（公式文本）：理解计算逻辑、追踪跨表引用
      · 通道 B（缓存值）  ：上一次 Excel 计算的结果
    程序生成、从未被 Excel 打开保存过的文件，**通道 B 会整片为空** ——
    本工具会**明确统计并告警**，不要把它当成 0。
    """
    import openpyxl

    try:
        p = _safe_path(path)
    except ValueError as e:
        return str(e)
    if not p.exists():
        return f"文件不存在：{path}"
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        return (f"read_excel 只支持 .xlsx / .xlsm（收到的是 {p.suffix or '无扩展名'}）。"
                f"旧版 .xls 请先用 Excel 另存为 .xlsx；CSV 请用 describe_csv。")

    try:
        wb_f = openpyxl.load_workbook(p, data_only=False)     # 通道 A：公式
        wb_v = openpyxl.load_workbook(p, data_only=True)      # 通道 B：缓存值
    except Exception as e:
        return f"打开工作簿失败：{type(e).__name__}: {e}"

    lines = [f"文件：{path}"
             f"（{p.stat().st_size / 1024:.1f} KB，工作表 {len(wb_f.sheetnames)} 张）",
             f"工作表：{' / '.join(wb_f.sheetnames)}"]

    for name in wb_f.sheetnames:
        wf, wv = wb_f[name], wb_v[name]
        n_row = wf.max_row or 0
        n_col = wf.max_column or 0
        lines.append("")
        lines.append(f"【{name}】{n_row} 行 × {n_col} 列")

        merged = [str(r) for r in wf.merged_cells.ranges][:10]
        lines.append(f"  合并单元格：{'、'.join(merged) if merged else '无'}")

        labels, numbers, formulas, cached_hit, cached_miss, units = [], [], [], 0, 0, []
        for row in wf.iter_rows():
            for c in row:
                fv = c.value
                if fv is None:
                    continue
                cord = c.coordinate
                if isinstance(fv, str) and fv.startswith("="):
                    formulas.append(f"{cord}: {fv}")
                    vv = wv[cord].value
                    if vv is None or (isinstance(vv, str) and vv.startswith("#")):
                        cached_miss += 1
                    else:
                        cached_hit += 1
                elif isinstance(fv, (int, float)):
                    numbers.append(f"{cord}={fv:g}")
                else:
                    txt = str(fv).strip()
                    if any(u in txt for u in ("单位", "万元", "百万元", "千元", "元)")):
                        units.append(f"{cord}=\"{txt[:24]}\"")
                    if c.column == 1 or len(txt) <= 12:
                        labels.append(f"{cord}={txt[:20]}")

        if units:
            lines.append(f"  ★ 单位线索：{' '.join(units[:6])}")
        if labels:
            lines.append(f"  标签（{len(labels)}）：{'  '.join(labels[:_EXCEL_MAX_LIST])}"
                         + (" …" if len(labels) > _EXCEL_MAX_LIST else ""))
        if numbers:
            lines.append(f"  数值（{len(numbers)}）：{'  '.join(numbers[:_EXCEL_MAX_LIST])}"
                         + (" …" if len(numbers) > _EXCEL_MAX_LIST else ""))
        if formulas:
            lines.append(f"  公式/通道A（{len(formulas)}）：{'  '.join(formulas[:_EXCEL_MAX_LIST])}"
                         + (" …" if len(formulas) > _EXCEL_MAX_LIST else ""))
            lines.append(f"  缓存值/通道B：{len(formulas)} 个公式格中 "
                         f"**{cached_miss} 个缺失、{cached_hit} 个有值**")
            if cached_miss and not cached_hit:
                lines.append("  ⚠️ 该表通道 B **完全不可用**（文件由程序生成、从未被 Excel 保存过）。"
                             "公式结果读不到，必须自己按通道A的公式求值，"
                             "并把「本表缓存值缺失」写进报告的「未覆盖声明」。")
            elif cached_miss:
                lines.append(f"  ⚠️ 有 {cached_miss} 个公式格缓存值缺失，读数时不要当成 0。")
        else:
            lines.append("  （本表没有公式，全是静态值）")

    lines.append("")
    lines.append("提示：跨表引用请用「表名!单元格」；报告里每个数字都要能指到这里列出的坐标。")
    return "\n".join(lines)


def _load_csv(path: str):
    """内部：安全读取项目内的 CSV。返回 (df, None) 或 (None, 错误信息)。"""
    import pandas as pd

    try:
        p = _safe_path(path)
    except ValueError as e:
        return None, str(e)
    if not p.exists():
        return None, f"文件不存在：{path}"
    try:
        return pd.read_csv(p), None
    except Exception as e:
        return None, f"读取失败：{type(e).__name__}: {e}"


def describe_csv(path: str) -> str:
    """工具：在**本地**分析 CSV 的结构与统计概况。

    只返回「列名 + 类型 + 缺失数 + 数值列统计」，**不返回任何明细行**。
    原因：工具结果会回传给大模型，明细里可能有姓名/学号等隐私，
    所以这里坚持"只出结构、不出明细"。
    """
    import pandas as pd

    df, err = _load_csv(path)
    if err:
        return err

    lines = [f"文件：{path}", f"行数：{len(df)}", f"列数：{len(df.columns)}", "各列概况："]
    for col in df.columns:
        s = df[col]
        miss = int(s.isna().sum())
        nums = pd.to_numeric(s, errors="coerce")
        if nums.notna().sum() >= max(1, len(s) // 2):        # 多数能转成数字 → 当数值列
            lines.append(f"  - {col}（数值列，缺失 {miss}）："
                         f"min={nums.min():g} max={nums.max():g} mean={nums.mean():.2f}")
        else:
            lines.append(f"  - {col}（文本列，缺失 {miss}）：不同取值 {s.nunique()} 个")
    return "\n".join(lines)


_AGGS = {"sum", "mean", "count", "max", "min"}
_IDENT = re.compile(r"[A-Za-z_\u4e00-\u9fff][A-Za-z0-9_\u4e00-\u9fff]*")


def _safe_expr(expr: str, columns) -> bool:
    """校验算式：只允许「列名 + 数字 + 运算符 + 括号」，防止执行任意代码。"""
    rest = _IDENT.sub("", expr)
    if not set(rest) <= set("0123456789+-*/(). "):
        return False
    return all(name in columns for name in _IDENT.findall(expr))


def aggregate_csv(path: str, group_by: str, value_expr: str = "",
                  agg: str = "sum", top_n: int = 10) -> str:
    """工具：在**本地**做分组聚合，只把「聚合结果」返回给模型。

    - group_by：按哪一列分组，如 "产品" / "地区" / "销售员"
    - value_expr：对哪一列（或算式，如 "数量*单价"）聚合；agg=count 时可留空
    - agg：sum / mean / count / max / min
    - top_n：只返回前几名（默认 10，最多 50）

    只返回聚合后的几行，**不返回任何明细行** —— 原始数据不出本机。
    """
    import pandas as pd

    df, err = _load_csv(path)
    if err:
        return err
    if group_by not in df.columns:
        return f"列名不存在：{group_by}。可用列：{list(df.columns)}"
    if agg not in _AGGS:
        return f"不支持的聚合方式：{agg}（可用：{sorted(_AGGS)}）"
    top_n = max(1, min(int(top_n), 50))

    if agg == "count" or not value_expr:
        series = df.groupby(group_by).size()
        label = "行数"
    else:
        if not _safe_expr(value_expr, df.columns):
            return (f"value_expr 只能用「列名 + 数字 + 运算符 + 括号」，不支持：{value_expr}")
        try:
            values = eval(value_expr, {"__builtins__": {}},
                          {c: df[c] for c in df.columns})    # noqa: S307（已用列名白名单校验）
        except Exception as e:
            return f"表达式计算失败：{type(e).__name__}: {e}"
        series = pd.Series(values).groupby(df[group_by]).agg(agg)
        label = f"{value_expr} 的 {agg}"

    series = series.sort_values(ascending=False).head(top_n)
    body = "\n".join(f"  {k}\t{v}" for k, v in series.items())
    return f"按「{group_by}」分组统计（{label}），前 {len(series)} 名：\n{body}"


# ---------------- 代码执行工具（数据分析主力）----------------
# ⚠️ 安全防线（粗粒度，**不是真沙箱**！）
# 这道 guard 只挡"最明显危险"的写法：装包、起子进程、删文件、访问项目外的路径。
# 黑名单永远可能被绕过，所以它只能算"第一道门"，真正的安全要靠进程/容器隔离。
_FORBIDDEN = [
    "subprocess", "os.system", "os.popen", "pip install", "pip3 install",
    "shutil.rmtree", "os.remove", "os.unlink", "os.rmdir", "os.rename", "shutil.move",
    "__import__", "importlib", "socket", "requests", "urllib", "httpx", "ftplib",
]
# 匹配代码里"带引号的绝对路径"，例如 r'C:\Users\...' 或 "D:/data/x.csv"
_ABS_PATH_TOKEN = re.compile(r"""['"]([A-Za-z]:[\\/][^'"]*)['"]""")


def _guard(code: str) -> str | None:
    """返回 None = 放行；返回字符串 = 拒绝执行的原因。"""
    low = code.lower()
    for bad in _FORBIDDEN:
        if bad in low:
            return f"拒绝执行：代码里出现了被禁止的用法「{bad}」（本项目只允许在 data/ 目录内做数据分析）。"

    for raw in _ABS_PATH_TOKEN.findall(code):
        try:
            p = Path(raw).resolve()
        except Exception:
            continue
        if p != BASE_DIR and BASE_DIR not in p.parents:
            return (f"拒绝执行：不允许访问项目目录之外的路径「{raw}」。"
                    f"请改用 DATA_DIR（它只指向本项目的 data/ 目录）。")
    return None


SANDBOX_DIR = BASE_DIR / "sandbox"      # 子进程的工作目录
MAX_OUTPUT_CHARS = 4000                 # 回传给模型的输出上限（防止海量明细出域）


def run_python(code: str, timeout: int = 30) -> str:
    """工具：在**独立子进程**里执行 Python 代码（可用 pandas / numpy），返回 print 的输出。

    相比"在主进程里 exec"，这里是纵深防御：
      1. `_guard` 预过滤：禁装包 / 禁子进程 / 禁访问项目外路径；
      2. **独立子进程**：代码崩了、死循环了都不影响主服务（有超时强杀）；
      3. **环境变量被精简**：子进程**拿不到任何 API key**，就算它想外传数据也没有凭据；
      4. **输出截断**：最多回传 MAX_OUTPUT_CHARS 字符，避免把明细批量灌回模型。

    ⚠️ 仍然不是"真沙箱"（Windows 上做不了只读挂载/禁网）。生产环境应该用容器
    或 gVisor 之类做进程级隔离。本工具**默认关闭**，需要时设 ENABLE_RUN_PYTHON=1 开启。
    """
    problem = _guard(code)
    if problem:
        return problem

    SANDBOX_DIR.mkdir(exist_ok=True)
    script = SANDBOX_DIR / "_tmp_run.py"
    prelude = ("import pandas as pd\n"
               "import numpy as np\n"
               f"DATA_DIR = r'{DATA_DIR}'\n")
    script.write_text(prelude + code, encoding="utf-8")

    # 精简环境变量：**故意不传** DEEPSEEK_API_KEY / SILICONFLOW_API_KEY
    safe_env = {
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "PATH": os.environ.get("PATH", ""),
        "SYSTEMROOT": os.environ.get("SYSTEMROOT", ""),
        "TEMP": os.environ.get("TEMP", ""),
        "TMP": os.environ.get("TMP", ""),
    }
    try:
        proc = subprocess.run(
            [sys.executable, str(script)],
            cwd=str(SANDBOX_DIR), env=safe_env,
            capture_output=True, text=True, encoding="utf-8", errors="replace",
            timeout=int(timeout),
        )
    except subprocess.TimeoutExpired:
        return f"执行超时（超过 {timeout} 秒），已强制终止。请把代码改得更简单些。"
    except Exception as e:
        return f"执行失败：{type(e).__name__}: {e}"

    out = (proc.stdout or "").strip()
    if proc.stderr:
        err = proc.stderr.strip()
        out = (out + "\n[stderr]\n" + err) if out else err
    if not out:
        out = "(代码执行完但没有输出，记得用 print 打印结果)"
    if len(out) > MAX_OUTPUT_CHARS:
        out = out[:MAX_OUTPUT_CHARS] + f"\n…（输出过长已截断，原始共 {len(out)} 字符）"

    # ★ 相对路径踩坑提示（放在截断**之后**追加，保证它一定不会被截掉）
    #   为什么需要：相对路径不会触发 _guard（它不是绝对路径），
    #   所以模型拿到的只有一坨原生 traceback，完全不知道
    #   是因为"代码跑在 sandbox/ 下、相对路径起点不对"。
    if "FileNotFoundError" in out or "No such file or directory" in out:
        out += (
            "\n\n[提示] 这个 FileNotFoundError 很可能是**相对路径**引起的："
            "本段代码的工作目录是 sandbox/，不是项目根目录，"
            "所以 'data/xxx.csv' 会被解析成 'sandbox/data/xxx.csv'（不存在）。"
            "读项目里的文件请用**绝对路径**：DATA_DIR + '/xxx.csv'，"
            "其中 DATA_DIR 已在本段代码开头自动定义好，指向本项目的 data/ 目录。"
        )
    return out
